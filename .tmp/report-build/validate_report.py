from pathlib import Path
from zipfile import ZipFile
from docx import Document
p = Path(__file__).resolve().parents[2] / 'deliverables/Bao_cao_du_an_QuanLyPharmacy_GPP.docx'
with ZipFile(p) as z:
    assert z.testzip() is None
d = Document(p)
assert len(d.tables) == 17
assert len(d.inline_shapes) == 1
assert len([x for x in d.paragraphs if x.style.name == 'Heading 1']) == 14
print('DOCX structure valid; 17 tables; 1 diagram; 14 chapter headings plus cover')
