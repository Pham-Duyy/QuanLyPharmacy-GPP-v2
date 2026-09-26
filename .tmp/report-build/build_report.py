from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'deliverables' / 'Bao_cao_du_an_QuanLyPharmacy_GPP.docx'
WORK = Path(__file__).parent
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.8)
sec.left_margin, sec.right_margin = Cm(2.3), Cm(2.0)
sec.header_distance, sec.footer_distance = Cm(.8), Cm(.8)
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(11.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.widow_control = True
for name, size in [('Title', 25), ('Heading 1', 17), ('Heading 2', 13), ('Heading 3', 12)]:
    st = doc.styles[name]
    st.font.name = 'Times New Roman'
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string('000000')
    st.font.bold = True
    st.paragraph_format.space_before = Pt(9)
    st.paragraph_format.space_after = Pt(7)
    st.paragraph_format.keep_with_next = True
doc.styles['Caption'].font.name = 'Times New Roman'
doc.styles['Caption'].font.size = Pt(10)
doc.styles['Caption'].font.color.rgb = RGBColor.from_string('444444')
sec.different_first_page_header_footer = True
foot = sec.footer.paragraphs[0]
foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = foot.add_run('QuanLyPharmacy GPP  |  ')
r.font.size = Pt(9)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); foot._p.append(fld)
doc.core_properties.title = 'Báo cáo dự án phần mềm quản lý nhà thuốc QuanLyPharmacy GPP'
doc.core_properties.subject = 'Báo cáo học thuật và tài liệu hỗ trợ thuyết trình'
doc.core_properties.author = 'Sinh viên thực hiện'
doc.core_properties.keywords = 'nhà thuốc, quản lý kho, FEFO, PostgreSQL, kiểm thử, nghiệp vụ'

def p(text, bold=False, style=None):
    a = doc.add_paragraph(style=style)
    a.add_run(text).bold = bold
    if not style: a.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return a

def h(text, level=2): return doc.add_heading(text, level)
def page(title):
    doc.add_page_break()
    h(title, 1)
def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    t.style = 'Table Grid'
    if widths is None: widths = [16.7 / len(headers)] * len(headers)
    for col, width in zip(t.columns, widths): col.width = Cm(width)
    for i, x in enumerate(headers): t.rows[0].cells[i].text = x
    for vals in rows:
        cells = t.add_row().cells
        for i, x in enumerate(vals): cells[i].text = str(x)
    for ri, row in enumerate(t.rows):
        for i, cell in enumerate(row.cells):
            cell.width = Cm(widths[i])
            for a in cell.paragraphs:
                a.paragraph_format.space_after = Pt(4)
                a.paragraph_format.space_before = Pt(3)
                a.paragraph_format.line_spacing = 1.03
                for r in a.runs:
                    r.font.name = 'Times New Roman'; r.font.size = Pt(10.5)
                    if ri == 0: r.bold = True
            if ri == 0:
                sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'E8EEF2'); cell._tc.get_or_add_tcPr().append(sh)
        trPr = row._tr.get_or_add_trPr()
        no = OxmlElement('w:cantSplit'); trPr.append(no)
        if ri == 0:
            repeat = OxmlElement('w:tblHeader'); trPr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def note(text):
    a = p(text)
    a.paragraph_format.left_indent = Cm(.3)
    for r in a.runs: r.italic = True

# Page 1
p('TRƯỜNG  ................................................................................', True).alignment = WD_ALIGN_PARAGRAPH.CENTER
p('KHOA  ....................................................................................').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph()
doc.add_paragraph()
p('BÁO CÁO DỰ ÁN', style='Title').alignment = WD_ALIGN_PARAGRAPH.CENTER
p('PHẦN MỀM QUẢN LÝ NHÀ THUỐC', style='Title').alignment = WD_ALIGN_PARAGRAPH.CENTER
p('QuanLyPharmacy GPP', True).alignment = WD_ALIGN_PARAGRAPH.CENTER
p('Quản lý nghiệp vụ bán lẻ theo lô và hạn dùng\nBảo vệ tính nhất quán dữ liệu khi nhiều quầy cùng thao tác').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph()
table(['Thông tin', 'Nội dung'], [
    ('Môn học', '........................................................................'),
    ('Sinh viên hoặc nhóm thực hiện', '........................................................................'),
    ('Mã số sinh viên và lớp', '........................................................................'),
    ('Giảng viên hướng dẫn', '........................................................................'),
    ('Ngày thuyết trình', '........................................................................'),
], [6,10.7])
p('TÓM TẮT', True)
p('Dự án xây dựng ứng dụng web quản lý nhà thuốc bán lẻ, kết nối danh mục, nhập hàng, bán hàng, đơn thuốc, tồn kho, trả hàng và báo cáo trong cùng hệ thống. Trọng tâm kỹ thuật là quản lý số lượng theo đơn vị nhỏ nhất, truy xuất đúng lô đã bán và bảo đảm các giao dịch đồng thời không làm sai tồn kho hoặc ghi nhận trùng chứng từ.')
p('Báo cáo trình bày bài toán, kiến trúc, mô hình dữ liệu, các quy trình nghiệp vụ, giải pháp xử lý giao dịch và kết quả kiểm thử. Phần cuối cung cấp kịch bản demo cùng lời dẫn cho bài thuyết trình khoảng 10 phút.')
note('Mốc tổng hợp nội dung: 26 tháng 9 năm 2026. Phạm vi là phần mềm hỗ trợ nghiệp vụ theo định hướng GPP; tên dự án không phải chứng nhận tuân thủ GPP hay đánh giá chuyên môn dược.')

# Page 2
page('Nội dung báo cáo')
table(['Phần', 'Nội dung', 'Trang'], [
('1', 'Bài toán và mục tiêu', '3'), ('2', 'Phạm vi chức năng và tác nhân', '4'),
('3', 'Kiến trúc và công nghệ', '5'), ('4', 'Thiết kế cơ sở dữ liệu', '6'),
('5', 'Nhập hàng và quản lý giá vốn', '7'), ('6', 'Bán hàng và xử lý đơn thuốc', '8'),
('7', 'Trả hàng kiểm kê và công nợ', '9'), ('8', 'Tính nhất quán và bảo mật', '10'),
('9', 'Kiểm thử và kết quả xác minh', '11'), ('10', 'Đánh giá giới hạn và phát triển', '12'),
('11', 'Kịch bản trình diễn', '13'), ('12', 'Lời dẫn thuyết trình 10 phút', '14'),
('13', 'Câu hỏi phản biện và nguồn đối chiếu', '15'),
], [1.3,13.8,1.6])
h('Thông điệp chính')
p('Giá trị của dự án nằm ở việc nối các nghiệp vụ thành một chuỗi có thể kiểm tra: hàng nhập tạo tồn và công nợ; bán hàng tạo hóa đơn, phân bổ lô và giá vốn; trả hàng hoặc hủy hóa đơn phải cập nhật đúng các dữ liệu liên quan. Một màn hình chạy được chưa đủ nếu số liệu giữa kho, chứng từ và báo cáo không khớp.')
h('Các điểm cần nhấn mạnh khi trình bày')
p('Thứ nhất, thiết kế theo đặc thù nhà thuốc: lô sản xuất, hạn dùng, đơn vị quy đổi, đơn thuốc và cảnh báo an toàn. Thứ hai, lựa chọn transaction, khóa hàng và khóa chống xử lý trùng để bảo vệ dữ liệu khi nhiều người thao tác. Thứ ba, kiểm thử các trường hợp lỗi và đối chiếu bằng kết quả kỳ vọng độc lập.')
h('Thuật ngữ dùng trong báo cáo')
table(['Thuật ngữ', 'Cách hiểu'], [
('FEFO', 'Ưu tiên xuất lô có hạn dùng gần nhất trong các lô đủ điều kiện bán.'),
('Transaction', 'Nhóm thay đổi dữ liệu được commit cùng nhau hoặc rollback cùng nhau.'),
('Idempotency', 'Gửi lại cùng một yêu cầu không tạo thêm một kết quả nghiệp vụ mới.'),
('Snapshot giá vốn', 'Lưu giá vốn tại thời điểm xuất để báo cáo không chạy theo giá vốn lô về sau.'),
], [4,12.7])

# Page 3
page('1 Bài toán và mục tiêu')
h('1 1 Bối cảnh bài toán')
p('Một sản phẩm tại nhà thuốc có thể được nhập nhiều lần, nằm ở nhiều lô và được bán theo hộp, vỉ hoặc viên. Mỗi lô có hạn dùng, giá vốn và trạng thái chất lượng riêng. Vì vậy, việc chỉ lưu một con số tồn tổng cho sản phẩm không đủ để chọn hàng khi bán, xác định lô cần thu hồi hoặc truy lại hàng khách mang trả.')
p('Khi có nhiều quầy, sai lệch còn có thể xuất hiện do thời điểm thao tác: hai người cùng bán lượng tồn cuối, hai yêu cầu cùng đổi điểm, hoặc một người hủy hóa đơn trong khi người khác nhận trả hàng. Nếu các bước chỉ kiểm tra tuần tự ở giao diện, cả hai yêu cầu có thể cùng thấy dữ liệu cũ và cùng được thực hiện.')
h('1 2 Vấn đề cần giải quyết')
table(['Vấn đề', 'Hệ quả', 'Hướng xử lý trong dự án'], [
('Quy đổi đơn vị', 'Bán một hộp nhưng trừ sai số viên.', 'Chuẩn hóa baseQuantity tại backend.'),
('Lô và hạn dùng', 'Xuất sai lô hoặc bán hàng không đủ điều kiện.', 'Lọc trạng thái và hạn dùng rồi phân bổ FEFO.'),
('Thao tác đồng thời', 'Bán vượt tồn hoặc hoàn tồn hai lần.', 'Khóa và kiểm tra trạng thái trong transaction.'),
('Lỗi mạng và gửi lại', 'Tạo thêm hóa đơn cho một lần mua.', 'Idempotency gắn với chứng từ đã commit.'),
('Nhập hàng khác giá', 'Sai giá vốn hoặc thay đổi lãi kỳ cũ.', 'Bình quân tồn và snapshot giá vốn xuất.'),
], [3.5,5.5,7.7])
h('1 3 Mục tiêu')
p('Mục tiêu nghiệp vụ là hỗ trợ quản lý vòng đời hàng hóa từ lúc nhập đến bán, trả, kiểm kê hoặc xử lý chất lượng; giảm thao tác đối chiếu thủ công và tăng khả năng truy xuất. Mục tiêu kỹ thuật là duy trì các ràng buộc dữ liệu, phân quyền theo cửa hàng và cung cấp bằng chứng kiểm thử cho những tình huống dễ gây sai lệch.')
h('1 4 Tiêu chí đánh giá')
p('Hệ thống được đánh giá qua tính đúng của quy trình, khả năng bảo toàn số lượng và giá trị, việc chặn thao tác không hợp lệ, khả năng phục hồi sau gửi lại yêu cầu và sự rõ ràng của báo cáo. Khả năng chịu tải lớn chưa được kết luận chỉ từ số lượng test đạt; cần đo riêng với dữ liệu và mức đồng thời đại diện.')
note('Cách giới thiệu đề tài: “Phần mềm quản lý nhà thuốc theo định hướng hỗ trợ thực hành GPP”, tránh khẳng định hệ thống đã được chứng nhận hoặc đã thay thế việc thẩm định chuyên môn.')

# Page 4
page('2 Phạm vi chức năng và tác nhân')
h('2 1 Nhóm chức năng đã triển khai')
table(['Nhóm', 'Nội dung tiêu biểu'], [
('Danh mục', 'Sản phẩm, nhóm hàng, hoạt chất, đơn vị quy đổi, mã vạch, hình ảnh, bảng giá và nhà cung cấp.'),
('Kho và mua hàng', 'Nhập hàng, tồn đầu kỳ, tồn theo lô, biệt trữ, kiểm kê, điều chỉnh, cảnh báo cận hạn và gợi ý bổ sung hàng.'),
('Bán hàng', 'Quầy bán, hóa đơn, giảm giá, tích và đổi điểm, phân bổ lô, in hóa đơn, trả hàng và hủy hóa đơn.'),
('Hồ sơ và chất lượng', 'Khách hàng, dị ứng, đơn thuốc và ảnh đơn, thuốc kiểm soát đặc biệt, thu hồi, nhật ký bảo quản.'),
('Tài chính vận hành', 'Công nợ nhà cung cấp, phân bổ thanh toán, trả hàng nhà cung cấp và báo cáo kinh doanh.'),
('Quản trị', 'Cửa hàng, người dùng, quyền, nhật ký thao tác, cấu hình, nhập/xuất Excel, tem nhãn và sao lưu.'),
], [4.1,12.6])
h('2 2 Tác nhân và phạm vi quyền')
table(['Vai trò', 'Trách nhiệm chính'], [
('Quản lý', 'Quản trị cửa hàng, người dùng, giá, báo cáo, công nợ và cấu hình theo quyền được cấp.'),
('Dược sĩ', 'Xác nhận đơn thuốc, xử lý cảnh báo cần xác nhận và các thao tác chuyên môn được phân quyền.'),
('Nhân viên bán hàng', 'Thực hiện bán hàng thông thường và quản lý thông tin khách trong phạm vi được cấp.'),
('Nhân viên kho', 'Lập phiếu nhập, kiểm kê, lập điều chỉnh và ghi nhận điều kiện bảo quản theo quyền.'),
('Người kiểm tra chỉ đọc', 'Tra cứu chứng từ, giá vốn, báo cáo và nhật ký; không mặc định được sửa dữ liệu.'),
], [4.1,12.6])
p('Vai trò được gán theo cửa hàng hoặc phạm vi toàn chuỗi. Quyền quản trị không được hiểu là tự động thay thế quyền chuyên môn của dược sĩ; khi demo cần dùng đúng tài khoản và tập quyền thực tế.')
h('2 3 Ranh giới phạm vi')
p('Thiết kế hỗ trợ nhiều cửa hàng nhưng không đồng nghĩa đã có đầy đủ nghiệp vụ điều chuyển liên chi nhánh. Đơn đặt mua chính thức, thanh toán qua cổng ngân hàng, hóa đơn điện tử, kiểm tra tương tác thuốc từ nguồn chuyên môn và vận hành ngoại tuyến không được trình bày như tính năng đã hoàn tất nếu chưa có triển khai và bằng chứng riêng.')

# Page 5
page('3 Kiến trúc và công nghệ')
p('Ứng dụng được tổ chức thành frontend React và backend Express, dữ liệu nghiệp vụ lưu tập trung trong PostgreSQL. Backend chia theo các mô đun chức năng trong cùng ứng dụng, phù hợp với quy mô hiện tại và giúp giữ các thay đổi kho, hóa đơn, điểm và nhật ký trong một transaction.')
table(['Lớp', 'Công nghệ và vai trò'], [
('Giao diện', 'React, TypeScript, Vite, Ant Design; React Query quản lý dữ liệu từ máy chủ; Axios gọi API.'),
('API và nghiệp vụ', 'Node.js, Express, TypeScript; Zod kiểm tra đầu vào; các service xử lý quy tắc nghiệp vụ.'),
('Truy cập dữ liệu', 'Prisma và PostgreSQL; SQL trực tiếp cho khóa hàng, tổng hợp và thao tác nguyên tử.'),
('Môi trường và chất lượng', 'Docker cho PostgreSQL; Vitest, Supertest; TypeScript và oxlint cho kiểm tra mã.'),
], [4,12.7])

im = Image.new('RGB', (1500,510), 'white'); d=ImageDraw.Draw(im)
font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 30)
small = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 25)
boxes = [(25,80,445,310),(540,80,960,310),(1055,80,1475,310)]
texts = [['GIAO DIỆN WEB','React và TypeScript','Quầy bán và quản trị'],['API VÀ NGHIỆP VỤ','Express và Prisma','Quyền  •  Giá  •  Tồn'],['DỮ LIỆU TẬP TRUNG','PostgreSQL','Lô  •  Chứng từ  •  Nhật ký']]
for box, lines in zip(boxes,texts):
    d.rounded_rectangle(box, radius=12, fill='#edf2f6', outline='#344c61', width=3)
    for j,line in enumerate(lines):
        x=(box[0]+box[2])/2
        d.text((x,box[1]+42+j*56),line,font=font if j==0 else small,fill='#162a3a',anchor='mm')
for a,b in [(445,540),(960,1055)]:
    d.line((a+5,195,b-13,195),fill='#344c61',width=4)
    d.polygon([(b-8,195),(b-23,186),(b-23,204)],fill='#344c61')
d.text((750,390),'Yêu cầu được kiểm tra tại backend trước khi thay đổi dữ liệu',font=font,fill='#162a3a',anchor='mm')
im.save(WORK/'architecture.png')
doc.add_picture(str(WORK/'architecture.png'), width=Cm(16.6))
p('Hình 1  Kiến trúc tổng quát và nơi quyết định nghiệp vụ', style='Caption')
h('3 1 Luồng xử lý một yêu cầu')
p('Giao diện gửi yêu cầu đến API /api/v1. Middleware xác thực tài khoản, xác định cửa hàng và kiểm tra quyền. Dữ liệu được kiểm tra bằng schema trước khi service thực thi. Những thao tác có ảnh hưởng đến tồn hoặc tiền được tổ chức trong transaction; kết quả trả về theo cấu trúc thống nhất để giao diện hiển thị và cập nhật dữ liệu.')
h('3 2 Lý do lựa chọn')
p('TypeScript giúp phát hiện lỗi kiểu dữ liệu sớm ở cả hai phía. PostgreSQL phù hợp với quan hệ chứng từ và ràng buộc chặt chẽ. Prisma thuận tiện cho truy vấn thông thường; SQL được dùng khi cần kiểm soát khóa và phép cập nhật nguyên tử. Chưa có nhu cầu đã chứng minh để tách hệ thống thành microservices.')

# Page 6
page('4 Thiết kế cơ sở dữ liệu')
p('Schema Prisma tại mốc khảo sát có 53 model. Đây là số cấu trúc dữ liệu được khai báo, không phải 53 tính năng đã hoàn thiện; ví dụ sự tồn tại của AiLog không chứng minh đã có mô hình AI hoạt động. Các nhóm dữ liệu chính được thiết kế quanh sản phẩm, lô hàng và chứng từ.')
table(['Nhóm dữ liệu', 'Bảng hoặc model tiêu biểu', 'Mục đích'], [
('Tài khoản và phạm vi', 'Store, User, Role, Permission, UserRole', 'Tách quyền thao tác và phạm vi cửa hàng.'),
('Danh mục', 'Product, ProductUnit, ProductPrice, ProductIngredient', 'Đơn vị quy đổi, giá theo thời gian và thành phần.'),
('Kho', 'Batch, StockMovement, GoodsReceipt, GoodsReceiptLine', 'Theo dõi lô, nhập kho và lịch sử biến động.'),
('Bán và trả', 'Invoice, InvoiceLine, InvoiceAllocation, Return, ReturnLine', 'Lưu hóa đơn, phân bổ lô và truy lại hàng trả.'),
('Chuyên môn', 'Prescription, PrescriptionItem, CustomerAllergy, Recall', 'Đơn thuốc, dị ứng và xử lý thu hồi.'),
('Kiểm soát vận hành', 'IdempotencyKey, DocumentCounter, AuditLog, Backup', 'Chống trùng, cấp mã, truy vết và sao lưu.'),
], [3.3,6.5,6.9])
h('4 1 Các quan hệ quan trọng')
p('Một sản phẩm có nhiều đơn vị và nhiều lô. Một hóa đơn có nhiều dòng hàng; mỗi dòng có thể được xuất từ nhiều lô thông qua InvoiceAllocation. Dòng trả hàng tham chiếu phân bổ gốc để xác định đúng lô đã xuất. Một phiếu nhập có nhiều dòng; xác nhận phiếu tạo hoặc tăng tồn ở lô tương ứng.')
h('4 2 Lựa chọn mô hình và ràng buộc')
p('Tồn được lưu theo đơn vị nhỏ nhất. Tiền giao dịch sử dụng số nguyên đồng; giá vốn trên đơn vị nhỏ nhất sử dụng số thập phân. Khóa ngoại bảo vệ liên kết, unique chống trùng định danh, CHECK bảo vệ các giới hạn dữ liệu. Những ràng buộc PostgreSQL không biểu diễn đủ trong Prisma được bổ sung bằng migration SQL.')
p('Dòng hóa đơn lưu tên, đơn vị, hệ số quy đổi và giá tại thời điểm bán. Phân bổ lưu giá vốn xuất cùng mức độ tin cậy ACTUAL, ESTIMATED hoặc UNKNOWN. Vì vậy, việc sửa danh mục hoặc nhập thêm với giá khác không tự động thay đổi ý nghĩa của giao dịch đã hoàn tất.')
h('4 3 Tồn hiện tại và thẻ kho')
p('Batch phục vụ đọc tồn nhanh, còn StockMovement ghi từng biến động. Việc duy trì đồng thời hai dạng dữ liệu đòi hỏi đối chiếu định kỳ: tồn đầu cộng nhập và hoàn, trừ bán và các lần xuất, phải khớp tồn cuối. Dự án có script kiểm tra các bất biến, trong đó có đối chiếu tồn với tổng biến động.')

# Page 7
page('5 Nhập hàng và quản lý giá vốn')
h('5 1 Quy trình nhập hàng')
table(['Bước', 'Xử lý và kiểm soát'], [
('Lập nháp', 'Chọn nhà cung cấp, sản phẩm, đơn vị, số lượng, giá nhập, số lô, hạn dùng và chiết khấu hoặc thuế.'),
('Kiểm tra', 'Đơn vị phải thuộc sản phẩm; quy đổi số lượng; cùng số lô phải nhất quán hạn dùng.'),
('Kiểm nhập', 'Ghi kết quả đạt hoặc không đạt cho từng dòng; hàng không đạt được biệt trữ.'),
('Xác nhận', 'Chuyển trạng thái có điều kiện; tạo hoặc cập nhật lô và ghi thẻ kho trong transaction.'),
('Theo dõi thanh toán', 'Phiếu mua đã xác nhận tham gia công nợ; hạn thanh toán được xác định theo kỳ hạn nhà cung cấp.'),
], [3.2,13.5])
h('5 2 Công thức giá vốn')
p('Khi nhập thêm cùng lô, hệ thống tính bình quân gia quyền trên lượng đang còn: giá vốn mới bằng tổng giá trị tồn trước nhập cộng giá trị đợt nhập, chia cho tổng số lượng sau nhập. Chiết khấu và thuế của phiếu được phân bổ vào giá vốn theo quy tắc của dự án.')
table(['Thành phần', 'Số lượng', 'Giá mỗi đơn vị', 'Giá trị'], [
('Tồn ban đầu', '100', '1.000', '100.000'),
('Nhập lần A', '100', '2.000', '200.000'),
('Nhập lần B', '100', '4.000', '400.000'),
('Sau hai lần nhập', '300', '2.333,3333', '700.000 trước sai số làm tròn'),
], [5.2,2.5,4.5,4.5])
p('Ví dụ trên vẫn phải đúng khi A và B xác nhận đồng thời. Nếu cả hai cùng đọc tồn ban đầu rồi ghi lại kết quả tính riêng, giá vốn có thể thành 2.500. Bản sửa tính số lượng và giá vốn ngay trong biểu thức UPDATE, để mỗi lần cập nhật dùng dữ liệu được bảo vệ bởi khóa hàng. Lô mới dùng INSERT có xử lý xung đột để không tạo hai bản ghi cùng lô.')
h('5 3 Giá vốn tồn khác giá vốn xuất')
p('Giá vốn tồn có thể thay đổi sau một đợt nhập. Giá vốn đã xuất phải được chụp trên phân bổ hóa đơn. Khi trả hàng RESTOCK hoặc hủy, lượng hàng hoàn được cộng lại theo giá vốn xuất gốc rồi tính lại bình quân tồn. Cách này giúp bảo toàn giá trị và giữ ổn định báo cáo của giao dịch trước.')
note('Các ví dụ tiền trong báo cáo là bộ số minh họa để giải thích thuật toán, không phải số liệu kinh doanh thật. Giá vốn được lưu với độ chính xác hữu hạn nên kiểm thử sử dụng dung sai phù hợp.')

# Page 8
page('6 Bán hàng và xử lý đơn thuốc')
h('6 1 Quy trình bán tại quầy')
p('Người bán chọn cửa hàng, tìm hoặc quét sản phẩm, chọn đơn vị và số lượng; nếu có thì chọn khách hàng và đơn thuốc. Giao diện hỗ trợ thao tác, nhưng backend mới là nơi quyết định quyền, giá hiện hành, mức giảm giá, lượng điểm được đổi và các lô đủ điều kiện xuất.')
table(['Thứ tự', 'Xử lý tại backend'], [
('1', 'Đọc sản phẩm và đơn vị; chuyển quantity thành baseQuantity.'),
('2', 'Kiểm tra quyền theo mặt hàng, giảm giá hoặc chỉ định lô; khóa dữ liệu cần thiết.'),
('3', 'Kiểm tra trạng thái và hiệu lực đơn thuốc, cảnh báo và xác nhận bắt buộc.'),
('4', 'Tính giá, phân bổ giảm giá, tách VAT và kiểm tra tiền thanh toán.'),
('5', 'Đối chiếu số đã cấp phát; khóa lô đủ điều kiện và phân bổ theo FEFO.'),
('6', 'Ghi hóa đơn, phân bổ, giá vốn xuất, trừ tồn, thẻ kho, điểm và nhật ký.'),
('7', 'Gắn chứng từ với khóa chống trùng rồi commit; giao diện nhận kết quả để in.'),
], [1.7,15])
h('6 2 Ví dụ quy đổi và chọn lô')
p('Giả sử một hộp có 10 vỉ, mỗi vỉ 10 viên. Bán 2 vỉ tương ứng 20 viên. Lô A còn 12 viên và hết hạn sớm hơn lô B. Nếu cả hai đủ điều kiện, hệ thống xuất 12 viên từ A và 8 viên từ B, lưu hai phân bổ cho cùng dòng hóa đơn. Hàng trả về sau phải xác định đúng phân bổ gốc.')
h('6 3 Đơn thuốc và cảnh báo')
p('Đối với mặt hàng được phân loại kê đơn hoặc kiểm soát đặc biệt, hệ thống kiểm tra quyền bán, yêu cầu đơn phù hợp, trạng thái xác nhận, hiệu lực và số lượng còn được cấp. Liên kết dòng đơn thuốc do backend xác định được lưu lại để hủy hoặc trả có thể cập nhật số đã cấp phát.')
p('Kiểm tra an toàn hiện dùng các quy tắc tất định, gồm dữ liệu dị ứng và trùng hoạt chất trong giỏ. Cảnh báo cần xác nhận phải được người đủ quyền ghi nhận. Những phần chưa có dữ liệu kiểm tra được biểu diễn là chưa kiểm tra, không coi việc không có cảnh báo là bằng chứng an toàn đầy đủ.')
note('Không giới thiệu tính năng này là AI chẩn đoán hoặc hệ thống kiểm tra tương tác thuốc toàn diện. Bộ quy tắc hỗ trợ thao tác và đối chiếu dữ liệu, không thay thế quyết định chuyên môn.')

# Page 9
page('7 Trả hàng kiểm kê và công nợ')
h('7 1 Trả hàng và hủy hóa đơn')
p('Trả hàng được lập từ hóa đơn gốc. Backend kiểm tra trạng thái hóa đơn, thời hạn nhận trả theo cấu hình, số đã trả và đúng đơn vị/lô. Tiền hoàn dựa trên số tiền thực tế đã lưu sau giảm giá, không dựa vào bảng giá hiện hành. Trả từng phần dùng tiền hoàn lũy kế để hạn chế sai số cộng dồn.')
p('Ví dụ dòng 6 đơn vị có tổng tiền 10.000 đồng: nếu mỗi lần trả một đơn vị đều làm tròn thành 1.667 đồng thì tổng sẽ thành 10.002 đồng. Hệ thống tính tiền hoàn lần này bằng tiền được hoàn theo tổng số đã trả mới trừ tiền theo tổng số đã trả trước đó. Khi trả hết, tổng hoàn khớp tiền dòng gốc.')
p('RESTOCK làm tăng số lượng và hoàn giá trị tồn. DISPOSE ghi nhận hàng vào rồi ra trong cùng giao dịch, không làm tăng tồn cuối. Hủy hóa đơn hoàn về các lô gốc, xử lý điểm và số đã cấp phát; không cho vừa hủy vừa nhận trả trên cùng hóa đơn theo các điều kiện nghiệp vụ.')
h('7 2 Kiểm kê và điều chỉnh')
p('Đợt kiểm kê cho phép ghi số đếm từng lô theo đơn vị thực tế. Hệ thống ghi mốc tồn lúc nhập số đếm, tính chênh lệch và khi chốt tạo phiếu điều chỉnh nháp. Người duyệt khác người lập theo quy tắc đã triển khai. Dòng chưa đếm không bị hiểu là bằng 0.')
p('Ví dụ tồn hệ thống lúc đếm là 100, thực đếm 98, chênh lệch là −2. Trước khi duyệt có bán thêm 10, tồn hệ thống còn 90. Duyệt điều chỉnh phải đưa tồn về 88, không gán lại bằng 98; nếu gán tuyệt đối sẽ làm mất ảnh hưởng của giao dịch bán xen kẽ.')
h('7 3 Công nợ và trả hàng nhà cung cấp')
p('Công nợ được xác định từ phiếu nhập đã xác nhận, trừ các khoản thanh toán còn hiệu lực và khoản giảm nợ hợp lệ. Khi phân bổ thanh toán, các phiếu nhập được khóa để tránh hai người cùng thanh toán trên số nợ cũ. Phiếu trả nhà cung cấp khi xác nhận trừ tồn và ghi thẻ kho.')
p('Đợt nghiệm thu gần nhất phát hiện trường hợp chọn trừ vào công nợ dù phiếu đã thanh toán đủ, làm mất dấu khoản cần được nhà cung cấp hoàn. Mã đang hoàn thiện đã bổ sung chặn DEBT_CREDIT_EXCEEDED và hướng dẫn chọn cách tất toán khác. Cần cập nhật kết quả nghiệm thu cuối cho bản sửa này trước khi tuyên bố hoàn tất.')

# Page 10
page('8 Tính nhất quán và bảo mật')
h('8 1 Bốn cơ chế bảo vệ dữ liệu')
table(['Cơ chế', 'Vai trò', 'Ví dụ'], [
('Transaction', 'Các thay đổi liên quan thành công hoặc thất bại cùng nhau.', 'Không để hóa đơn đã tạo nhưng tồn chưa trừ.'),
('Khóa hàng', 'Kiểm tra và cập nhật trên dữ liệu được bảo vệ khi tranh chấp.', 'Hai quầy cùng bán phần tồn cuối.'),
('Chuyển trạng thái có điều kiện', 'Chỉ thao tác khi chứng từ còn ở trạng thái hợp lệ.', 'Phiếu nháp chỉ xác nhận một lần.'),
('Idempotency và quyền sở hữu', 'Gửi lại không tạo nghiệp vụ trùng; chủ cũ không commit sau tiếp quản.', 'Mất response sau khi hóa đơn đã commit.'),
], [4.4,6.3,6])
p('Khóa idempotency lưu chữ ký yêu cầu, người dùng, phạm vi cửa hàng và chứng từ đã tạo. Gắn chứng từ được thực hiện trong cùng transaction với nghiệp vụ. Khi khóa treo được tiếp quản, ownerToken thay đổi. Nếu request cũ không cập nhật được đúng một khóa đang xử lý thuộc về mình, service ném lỗi để rollback.')
p('Mã chứng từ sử dụng bảng bộ đếm tăng nguyên tử thay cho COUNT + 1. Khóa duy nhất vẫn được giữ ở database. Với nhiều loại nghiệp vụ cùng chạm dữ liệu, thứ tự khóa và thời gian giữ khóa cần được rà soát và đo, thay vì coi việc đã có transaction là đủ để loại bỏ mọi deadlock.')
h('8 2 Xác thực và phân quyền')
p('Access token được giữ trong bộ nhớ frontend; refresh token dùng cookie HttpOnly và có cơ chế xoay vòng. Mật khẩu được băm. Backend kiểm tra permission và phạm vi cửa hàng của yêu cầu; việc ẩn nút ở giao diện chỉ hỗ trợ trải nghiệm, không phải hàng rào bảo mật chính.')
p('Hệ thống tách quyền xem giá vốn và dữ liệu khách hàng nhạy cảm. Giao dịch quan trọng có nhật ký để truy người thực hiện và chứng từ liên quan. Các cơ chế này là lớp bảo vệ đã triển khai, không thay thế kiểm thử bảo mật và đánh giá cấu hình khi triển khai thực tế.')
h('8 3 Đổi cửa hàng và phục hồi thao tác')
p('Mã cập nhật trong đợt nghiệm thu xử lý xóa cache khi chuyển cửa hàng, ghi nhớ lựa chọn cửa hàng và kiểm tra quyền khi khôi phục. Quầy bán cũng bổ sung xử lý REQUEST_ALREADY_COMMITTED để mở chứng từ đã lập thay vì tạo lại. Các thay đổi này cần được đưa vào bản nghiệm thu cuối cùng và kiểm tra cả request trả về muộn, giỏ hàng và quyền đã bị thu hồi.')

# Page 11
page('9 Kiểm thử và kết quả xác minh')
h('9 1 Phương pháp kiểm thử')
p('Kiểm thử tích hợp chạy trên PostgreSQL riêng để kiểm tra cả API, transaction và ràng buộc database. Vitest tổ chức ca kiểm thử; Supertest gọi API. Các test dùng chung database chạy tuần tự để tránh xóa dữ liệu của nhau. Test đồng thời có thể điều phối khóa để bảo đảm thực sự xảy ra chồng lấn thay vì chỉ gọi hai hàm gần nhau.')
p('Bộ số kỳ vọng được tính độc lập với hàm cần kiểm tra: số lượng tồn, giá trị tồn, tiền hoàn và số chứng từ được đối chiếu sau giao dịch. Các test review đã phát hiện lỗi thực tế rồi được giữ làm test hồi quy, không chỉ kiểm tra đường đi thành công.')
table(['Nhóm kiểm tra', 'Điều cần chứng minh'], [
('Kho và giá vốn', 'Không âm tồn; quy đổi đúng; nhập đồng thời không ghi đè giá vốn; giá trị tồn được hoàn đúng.'),
('Bán và trả', 'Chọn đúng lô; không hoàn tồn hai lần; tiền hoàn tổng khớp hóa đơn; cấp phát không vượt.'),
('Chống trùng', 'Retry không tạo thêm chứng từ; mất quyền sở hữu phải rollback; khóa mới không bị chủ cũ sửa.'),
('Báo cáo', 'Giữ giá vốn lịch sử; ghi rõ dữ liệu ước tính/không xác định; xét hàng trả từ kỳ trước.'),
('Phạm vi và giao diện', 'Không lẫn dữ liệu cửa hàng; đúng quyền; phục hồi sau lỗi mạng hoặc tải lại trang.'),
], [4.2,12.5])
h('9 2 Kết quả có thể trình bày')
table(['Mốc và nguồn kết quả', 'Trạng thái'], [
('Xác minh độc lập tại commit 93b6e96', '22/22 test liên quan đạt: 2 test review, 13 test idempotency, 7 test chất lượng giá vốn.'),
('Kiểm tra mã và đóng gói tại mốc trên', 'Backend typecheck/lint đạt; frontend build đạt, còn cảnh báo chunk lớn.'),
('Toàn bộ bộ test tại mốc báo cáo trước', 'Nhóm triển khai báo cáo 525/525 đạt; đây không phải kết quả full suite được xác minh độc lập trong lần review cuối.'),
('Đợt nghiệm thu đang tiếp tục', 'Có thêm test và thay đổi chưa commit; kết quả full suite cuối của đợt này chưa được cung cấp tại thời điểm tổng hợp.'),
], [6.1,10.6])
note('Khi bảo vệ, gắn mọi con số test với commit và lần chạy. Không dùng “525 test đạt” để khẳng định bản mã mới nhất đã được nghiệm thu toàn bộ hoặc hệ thống không còn lỗi.')

# Page 12
page('10 Đánh giá giới hạn và phát triển')
h('10 1 Kết quả đạt được')
p('Dự án đã hình thành chuỗi nghiệp vụ tương đối đầy đủ cho nhà thuốc bán lẻ, từ danh mục đến nhập, bán, trả, kiểm kê và báo cáo. Các điểm nổi bật là quản lý tồn theo lô, quy đổi đơn vị, snapshot giá bán và giá vốn, phân quyền cửa hàng, cùng cơ chế phục hồi yêu cầu đã commit.')
p('Quá trình review cho thấy giá trị của kiểm thử theo tình huống: lỗi hoàn tồn đồng thời, giá vốn nhập cùng lô, làm tròn tiền hoàn và thiếu cảnh báo giá vốn đều có thể tồn tại dù thao tác đơn lẻ thông thường chạy được. Việc giữ test tái hiện giúp kiểm soát các lần thay đổi sau.')
h('10 2 Hạn chế cần công khai')
table(['Hạn chế', 'Ý nghĩa và hướng xử lý'], [
('Giá vốn lịch sử', 'ESTIMATED là ước tính đóng băng khi chuyển đổi; UNKNOWN chưa xác định. Báo cáo phải hiển thị chất lượng dữ liệu.'),
('Kiểm tra thuốc', 'Chưa có nguồn tương tác thuốc toàn diện; không xem bộ quy tắc hiện tại như AI chẩn đoán.'),
('Hiệu năng', 'Chưa có số đo tải chứng minh năng lực vận hành quy mô lớn; cần đo thời gian phản hồi và chờ khóa.'),
('Triển khai và dự phòng', 'Có chức năng sao lưu nhưng vẫn cần diễn tập khôi phục, kiểm tra lịch, tính đầy đủ của ảnh và dữ liệu.'),
('Tự động hóa chất lượng', 'Repository có cấu hình CI nhưng hiện ở file ci.yml.disabled; không trình bày như pipeline đang chạy tự động.'),
('Nghiệm thu cuối', 'Còn thay đổi của đợt nghiệm thu chưa commit; cần chốt phiên bản và lưu bằng chứng kiểm tra.'),
], [4.6,12.1])
h('10 3 Lộ trình đề xuất')
p('Ưu tiên gần nhất là hoàn tất nghiệm thu theo một phiên bản cố định, bổ sung các tình huống hai quầy và chạy thử có giám sát. Tiếp theo là đo tải với dữ liệu đại diện, tối ưu truy vấn báo cáo và thời gian giữ khóa, đồng thời khôi phục CI và diễn tập phục hồi dữ liệu.')
p('Các hướng mở rộng gồm điều chuyển hàng giữa cửa hàng, đơn đặt mua, tích hợp thanh toán và hóa đơn điện tử, hoặc nguồn dữ liệu chuyên môn được thẩm định. Mỗi hướng cần yêu cầu nghiệp vụ, nguồn dữ liệu và tiêu chí nghiệm thu riêng; không nên đánh đổi độ đúng của luồng cốt lõi để tăng số lượng tính năng.')
h('10 4 Kết luận dự án')
p('Hệ thống có nền tảng phù hợp để tiếp tục chạy thử và hoàn thiện. Đóng góp chính của đề tài là mô hình hóa đúng quan hệ giữa lô, chứng từ và giá trị, đồng thời minh họa cách dùng giao dịch và kiểm thử để bảo vệ nghiệp vụ. Kết luận sẵn sàng dùng chính thức cần dựa trên nghiệm thu cuối, vận hành thử và yêu cầu triển khai thực tế.')

# Page 13
page('11 Kịch bản trình diễn')
p('Chuẩn bị dữ liệu demo riêng, không dùng thông tin khách hàng thật. Chọn một sản phẩm bán thông thường để minh họa nhanh và một sản phẩm kê đơn cho phần kiểm soát. Hai tài khoản hoặc hai phiên trình duyệt phải được đăng nhập trước để tránh mất thời gian trên lớp.')
table(['Thời lượng', 'Thao tác', 'Điều cần giải thích'], [
('30 giây', 'Đăng nhập và chỉ ra cửa hàng đang chọn.', 'Tài khoản, quyền và phạm vi cửa hàng khác nhau.'),
('45 giây', 'Mở sản phẩm có đơn vị hộp/vỉ/viên và hai lô.', 'Tồn được quản lý theo đơn vị nhỏ nhất; mỗi lô có hạn riêng.'),
('60 giây', 'Lập hoặc mở phiếu nhập, xác nhận và xem thẻ kho.', 'Nháp chưa tăng tồn; xác nhận cập nhật dữ liệu trong transaction.'),
('60 giây', 'Bán số lượng đi qua hai lô, mở chi tiết hóa đơn.', 'FEFO và InvoiceAllocation giải thích hàng xuất từ đâu.'),
('45 giây', 'Nhận trả một phần, xem tồn và tiền hoàn.', 'Giá bán và giá vốn lấy từ giao dịch gốc; không lấy giá mới.'),
('45 giây', 'Mở báo cáo hoặc ca kiểm thử hai quầy đã chuẩn bị.', 'Chứng minh số liệu nhất quán và giới hạn cảnh báo.'),
], [2.3,7.4,7])
h('Ví dụ số liệu dễ đối chiếu')
p('Dùng 100 viên giá vốn 1.000 đồng, giá bán 3.000 đồng, bán 10 viên. Khi chưa có giảm giá hoặc trả hàng, doanh thu là 30.000 đồng, giá vốn là 10.000 đồng, lãi gộp theo công thức dự án là 20.000 đồng. Nhập thêm cùng lô 100 viên giá 3.000 đồng; báo cáo của lần bán cũ phải giữ nguyên.')
p('Sau đó trả RESTOCK 4 viên: còn 6 viên bán ròng, doanh thu còn 18.000 đồng và giá vốn bán ròng còn 6.000 đồng; lãi gộp tương ứng là 12.000 đồng. Phân biệt đây là ví dụ tính theo quy tắc báo cáo dự án, không phải báo cáo kế toán tài chính đầy đủ.')
h('Phương án dự phòng')
p('Chuẩn bị trước ảnh các màn hình chính, kết quả test và một hóa đơn mẫu. Nếu máy chủ hoặc Docker gặp lỗi trong buổi báo cáo, chuyển sang giải thích chuỗi dữ liệu và bằng chứng đã lưu; không thực hiện reset database ngay trước giảng viên. Dành phần kiểm thử đồng thời cho video hoặc kết quả có thể lặp lại nếu thời gian demo trực tiếp hạn chế.')
h('Điểm kiểm tra trước buổi báo cáo')
p('Chốt commit; xác nhận database demo; kiểm tra quyền tài khoản, hạn dùng mẫu, tồn đầu và bảng giá; thử in hoặc xem trước hóa đơn; tắt thông báo riêng tư; mở sẵn báo cáo và file kết quả test. Với bài nói 10 phút, chỉ chọn 2–3 thao tác tiêu biểu, phần còn lại dùng hình hoặc giải thích.')

# Page 14
page('12 Lời dẫn thuyết trình 10 phút')
h('Phút 0 đến 1 Giới thiệu bài toán')
p('Kính thưa thầy cô, em xin trình bày dự án QuanLyPharmacy GPP, phần mềm quản lý nhà thuốc bán lẻ. Điểm đặc thù của bài toán là một sản phẩm có thể có nhiều lô, nhiều hạn dùng và được bán theo nhiều đơn vị. Vì vậy, phần mềm cần trả lời không chỉ còn bao nhiêu hàng, mà còn hàng thuộc lô nào, có được phép bán và khi khách trả thì phải về đâu.')
h('Phút 1 đến 2 Mục tiêu và phạm vi')
p('Dự án kết nối các nghiệp vụ danh mục, nhập hàng, bán hàng, đơn thuốc, trả hàng, kiểm kê, công nợ và báo cáo. Hệ thống có phân quyền theo vai trò và cửa hàng. Em tập trung vào tính chính xác của dữ liệu và khả năng truy xuất, đặc biệt trong trường hợp hai quầy thao tác cùng lúc. Phần kiểm tra an toàn hiện là quy tắc từ dữ liệu, không phải AI chẩn đoán.')
h('Phút 2 đến 3 Kiến trúc và dữ liệu')
p('Giao diện sử dụng React và TypeScript; backend sử dụng Express, Prisma và PostgreSQL. Backend chịu trách nhiệm quyết định quyền, giá và tồn. Quan hệ quan trọng nhất là hóa đơn, dòng hóa đơn và phân bổ lô: một dòng hàng có thể lấy từ nhiều lô, và dữ liệu này giúp truy lại chính xác khi trả hoặc hủy.')
h('Phút 3 đến 5 Minh họa nghiệp vụ chính')
p('Em xin minh họa một sản phẩm có đơn vị hộp, vỉ và viên. Hệ thống quy đổi về viên trước khi tính tồn. Khi bán, các lô đủ điều kiện được chọn theo hạn dùng gần nhất. Sau khi thanh toán, hóa đơn lưu các lô đã xuất và giá vốn tại thời điểm bán. Nhập thêm cùng lô với giá khác chỉ thay đổi giá vốn tồn, không được thay đổi lãi của hóa đơn cũ.')
h('Phút 5 đến 7 Giải pháp kỹ thuật nổi bật')
p('Để tránh sai dữ liệu khi nhiều người cùng thao tác, dự án kết hợp transaction, khóa hàng và chuyển trạng thái có điều kiện. Khi mất mạng, khóa idempotency giúp lần gửi lại không tạo hóa đơn thứ hai. Nếu khóa đã được request khác tiếp quản, request cũ phải rollback. Với nhập hàng, phép tính bình quân được thực hiện ngay trong lệnh cập nhật để tránh hai giao dịch cùng tính trên số tồn cũ.')
h('Phút 7 đến 9 Kiểm thử và giới hạn')
p('Các test sử dụng PostgreSQL thật và đối chiếu số lượng, giá trị, số chứng từ sau giao dịch. Những lỗi đã được tái hiện gồm hoàn tồn hai lần, giá vốn bị ghi đè khi nhập đồng thời và tiền hoàn bị làm tròn vượt tổng. Ở mốc review 93b6e96, 22 test trọng điểm đã được xác minh độc lập và đều đạt. Kết quả full suite cần nêu cùng phiên bản và lần chạy. Hệ thống vẫn cần đo tải, diễn tập khôi phục và hoàn tất nghiệm thu trước khi vận hành chính thức.')
h('Phút 9 đến 10 Kết luận')
p('Qua dự án, em thể hiện việc phân tích nghiệp vụ và chuyển quy tắc thành mô hình dữ liệu, transaction và kiểm thử. Giá trị chính là tính nhất quán từ nhập, bán đến trả và báo cáo. Hướng tiếp theo là chạy thử có giám sát, tối ưu theo số đo và mở rộng các tích hợp sau khi luồng cốt lõi đã ổn định. Em xin cảm ơn thầy cô và sẵn sàng trả lời câu hỏi.')

# Page 15
page('13 Câu hỏi phản biện và nguồn đối chiếu')
table(['Câu hỏi có thể gặp', 'Gợi ý trả lời'], [
('Vì sao quản lý theo lô thay vì tồn tổng?', 'Để phân biệt hạn dùng, trạng thái và giá vốn; đồng thời truy xuất đúng hàng đã bán, trả hoặc thu hồi.'),
('FEFO khác FIFO thế nào?', 'FEFO ưu tiên hạn dùng gần nhất; FIFO ưu tiên thứ tự nhập. Dự án chọn FEFO trong tập lô đủ điều kiện bán.'),
('Có transaction rồi vì sao còn cần khóa?', 'Transaction bảo vệ tính nguyên tử; khóa hoặc cập nhật có điều kiện bảo vệ quyết định dựa trên dữ liệu khi nhiều giao dịch tranh chấp.'),
('Idempotency khác unique mã hóa đơn?', 'Unique chặn hai bản ghi trùng mã; idempotency nhận biết hai request là cùng một ý định nghiệp vụ dù có thể sinh mã khác.'),
('Vì sao lưu giá vốn trên phân bổ?', 'Giá vốn lô thay đổi khi nhập thêm. Snapshot giữ giá vốn xuất gốc để báo cáo lịch sử và hoàn hàng không chạy theo giá mới.'),
('Làm sao biết báo cáo chưa đủ chính xác?', 'API trả chất lượng giá vốn. ESTIMATED và UNKNOWN được cảnh báo; so sánh lợi nhuận chỉ hiển thị như số đầy đủ khi hai kỳ đủ dữ liệu.'),
('Dự án đã đạt chuẩn GPP hay có AI chưa?', 'Dự án hỗ trợ các nghiệp vụ theo định hướng GPP. Chưa dùng tên đề tài để khẳng định chứng nhận; kiểm tra an toàn hiện là quy tắc tất định.'),
('525 test đạt có nghĩa không còn lỗi?', 'Không. Test chứng minh các kịch bản đã kiểm tra tại một phiên bản; vẫn cần kiểm thử bảo mật, tải và chạy thử thực tế.'),
], [6.1,10.6])
h('Nguồn đối chiếu trong dự án')
p('[1] README.md và package.json của client/server: cấu trúc, công nghệ và cách chạy. [2] docs/api-contract.md: quy tắc nghiệp vụ và API; docs/erd.md: mô hình thiết kế ban đầu. [3] server/prisma/schema.prisma cùng các migration: mô hình triển khai hiện tại. Tài liệu cũ có thể chưa phản ánh đầy đủ số model mới.')
p('[4] Các mô đun inventory, sales, loyalty, reports, auth và middlewares: thuật toán, quyền, giao dịch và idempotency. [5] Các test costing-value, codex-review-regressions, codex-followup-review, idempotency và cost-quality: bằng chứng hồi quy. [6] server/src/acceptance: các kịch bản nghiệm thu đang bổ sung.')
p('[7] Commit 93b6e96 và kết quả review độc lập 22/22 test; báo cáo nhóm triển khai về 525/525 test thuộc lần chạy đã nêu. [8] .github/workflows/ci.yml.disabled: trạng thái cấu hình CI tại thời điểm khảo sát.')
note('Trước khi nộp: điền thông tin bìa, chốt commit trình diễn và cập nhật kết quả full suite cuối. Nếu sử dụng công cụ AI để hỗ trợ lập trình hoặc soạn báo cáo, trình bày trung thực theo quy định môn học và bảo đảm tự giải thích được các quyết định kỹ thuật.')

# Apply language and table paragraph controls.
for a in doc.paragraphs:
    for r in a.runs:
        rp = r._element.get_or_add_rPr()
        lang = OxmlElement('w:lang'); lang.set(qn('w:val'), 'vi-VN'); rp.append(lang)
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(str(OUT))
print('paragraphs', len(doc.paragraphs), 'tables', len(doc.tables), 'explicit pages', 15)
